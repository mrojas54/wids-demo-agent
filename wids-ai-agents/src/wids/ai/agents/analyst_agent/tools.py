import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import classification_report, confusion_matrix, accuracy_score
from langchain.tools import tool


@tool
def perform_eda() -> str:
    """Perform exploratory data analysis on the diabetes dataset."""

    df = pd.read_csv(
        "/Users/michellerojas/pydata/virginia2025/wids-demo-agent/wids-ai-agents/src/wids/ai/agents/analyst_agent/diabetes.csv"
    )

    # Basic statistics
    stats = df.describe().to_dict()

    # Check for missing values
    missing_values = df.isnull().sum().to_dict()

    # Check for zero values (which might be missing values in this dataset)
    zero_values = {col: (df[col] == 0).sum() for col in df.columns}

    # Correlation with target
    correlations = df.corr()["Outcome"].drop("Outcome").to_dict()

    # Result dictionary
    eda_results = {
        "basic_stats": stats,
        "missing_values": missing_values,
        "zero_values": zero_values,
        "correlations_with_target": correlations,
    }
    return eda_results


@tool
def train_model():
    """Train a Random Forest model on the diabetes dataset."""

    df = pd.read_csv(
        "/Users/michellerojas/pydata/virginia2025/wids-demo-agent/wids-ai-agents/src/wids/ai/agents/analyst_agent/diabetes.csv"
    )

    # Replace zeros with NaN for certain columns where zero doesn't make biological sense
    cols_with_zeros = ["Glucose", "BloodPressure", "BMI"]
    for col in cols_with_zeros:
        df[col] = df[col].replace(0, np.nan)

        # Simple imputation with median
    for col in cols_with_zeros:
        df[col] = df[col].fillna(df[col].median())

        # Split features and target
    X = df.drop("Outcome", axis=1)
    y = df["Outcome"]

    # Train-test split
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=42, stratify=y
    )

    # Train model
    model = RandomForestClassifier(
        n_estimators=3,
        max_depth=3,
        min_samples_split=20,
        max_features="sqrt",
        class_weight="balanced",
        bootstrap=True,
        random_state=42,
        max_samples=0.8,
        min_samples_leaf=10,
    )
    model.fit(X_train, y_train)

    # Evaluate
    y_pred = model.predict(X_test)
    accuracy = accuracy_score(y_test, y_pred)
    class_report = classification_report(y_test, y_pred, output_dict=True)
    conf_matrix = confusion_matrix(y_test, y_pred).tolist()

    # Feature importance
    feature_importances = dict(zip(X.columns, model.feature_importances_))

    # Results
    model_results = {
        "accuracy": accuracy,
        "classification_report": class_report,
        "confusion_matrix": conf_matrix,
        "feature_importances": feature_importances,
    }

    return model_results


@tool
def presentation(content: str, theme_path: str = None):
    """
    Build a PowerPoint presentation with a background theme.

    Args:
        content: Text content to include in the presentation.
        theme_path: (Optional) Path to a PowerPoint theme/template (.pptx) file.
    """
    try:
        import os
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        # Load a presentation with a theme if provided, else create a new one
        if theme_path and os.path.exists(theme_path):
            prs = Presentation(theme_path)
        else:
            prs = Presentation()

        def add_content_slide(prs, title_text, content_text):
            """Adds a slide with the given title and content while keeping the theme formatting."""
            slide_layout = prs.slide_layouts[1]  # Title and Content Layout
            slide = prs.slides.add_slide(slide_layout)

            title = slide.shapes.title
            title.text = title_text
            title.text_frame.paragraphs[0].font.bold = True
            title.text_frame.paragraphs[0].font.size = Pt(32)

            # Add text content
            text_box = slide.shapes.add_textbox(
                Inches(1), Inches(1.5), Inches(8), Inches(5)
            )
            text_frame = text_box.text_frame
            text_frame.word_wrap = True
            text_box.line.fill.background()  # Removes outline

            for paragraph in content_text.split("\n"):
                p = text_frame.add_paragraph()
                p.text = paragraph
                p.font.size = Pt(20)

        # Add a title slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = "Diabetes Prediction"
        slide.placeholders[1].text = "Agent Workflow"

        # Split content into smaller parts
        max_chars_per_slide = 500
        chunks = [
            content[i : i + max_chars_per_slide]
            for i in range(0, len(content), max_chars_per_slide)
        ]

        # Add slides dynamically
        for i, chunk in enumerate(chunks):
            add_content_slide(prs, f"About Data (Part {i+1})", chunk)

        # Save presentation
        file_path = os.path.abspath("diabetes_presentation.pptx")
        prs.save(file_path)

        return {
            "status": "success",
            "message": f"Presentation created successfully at {file_path}",
            "file_path": file_path,
        }

    except Exception as e:
        return {
            "status": "error",
            "message": f"Failed to create presentation: {str(e)}",
        }
